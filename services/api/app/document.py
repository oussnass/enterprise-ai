from pathlib import Path
import re
from pypdf import PdfReader
from docx import Document as DocxDocument
import unicodedata

def sanitize_text(value: str) -> str:
    value=value.encode('utf-8', errors='replace').decode('utf-8')
    value=unicodedata.normalize('NFC', value)
    return ''.join(char for char in value if char in '\n\r\t' or ord(char) >= 32)

def ocr_pdf(data: bytes) -> str:
    import pytesseract
    from pdf2image import convert_from_bytes
    pages=convert_from_bytes(data, dpi=250, fmt='png', thread_count=1)
    return sanitize_text('\n\n'.join(pytesseract.image_to_string(page, config='--psm 6') for page in pages))

def extract_text(filename: str, data: bytes) -> str:
    ext=Path(filename).suffix.lower()
    if ext=='.pdf':
        import io
        reader=PdfReader(io.BytesIO(data))
        text=sanitize_text('\n\n'.join((p.extract_text(extraction_mode='layout') or '') for p in reader.pages)).strip()
        return text or ocr_pdf(data)
    if ext=='.docx':
        import io
        d=DocxDocument(io.BytesIO(data)); parts=[p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            parts.append('\n'.join(' | '.join(cell.text.strip() for cell in row.cells) for row in table.rows))
        return '\n\n'.join(parts)
    if ext in {'.txt','.md','.csv','.json','.xml','.html'}:
        return sanitize_text(data.decode('utf-8', errors='ignore'))
    if ext in {'.png','.jpg','.jpeg','.tif','.tiff','.webp'}:
        import io
        from PIL import Image
        import pytesseract
        return sanitize_text(pytesseract.image_to_string(Image.open(io.BytesIO(data)), config='--psm 6'))
    raise ValueError(f'Unsupported document type: {ext}')

def chunk_text(text: str, size=1000, overlap=150):
    text='\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in text.splitlines())
    text=re.sub(r'\n{3,}', '\n\n', text).strip()
    chunks=[]; start=0
    while start < len(text):
        end=min(len(text), start+size); chunks.append(text[start:end])
        if end==len(text): break
        start=max(0,end-overlap)
    return chunks

def salary_scale_markdown_from_text(text: str) -> str | None:
    def normalize_numbers(value: str) -> list[str]:
        numbers=[]
        for token in re.findall(r'\d[\d ,.]*\d|\d+', value):
            digits=re.sub(r'\D','',token)
            if len(digits) >= 5:
                numbers.append(f'{int(digits):,}'.replace(',', ' '))
        return numbers

    rows={}; current=None
    for line in text.splitlines():
        match=re.match(r'^\s*(4|5|6|7|8|9|10|11|12)\b(.*)$',line)
        if match:
            values=normalize_numbers(match.group(2))
            if len(values) >= 3:
                current=match.group(1); rows[current]=values
            continue
        if current:
            rows[current].extend(normalize_numbers(line))
    if not rows:
        section=text.lower().split('salary scale',1)[-1] if 'salary scale' in text.lower() else ''
        fallback=[]
        for line in section.splitlines():
            values=re.findall(r'\d{4,6}',line)
            if len(values) >= 3:
                fallback.append(values)
            elif fallback and values:
                fallback[-1].extend(values)
        rows={str(index + 4): values for index,values in enumerate(fallback[:9])}
    if not rows:
        return None
    output=['| Echelon | Valeurs salariales extraites |','| --- | --- |']
    output.extend(f"| {echelon} | {' '.join(rows[echelon])} |" for echelon in sorted(rows,key=int))
    return '\n'.join(output)

def extract_salary_scale_markdown(data: bytes) -> str | None:
    import io

    def normalize_numbers(text: str) -> list[str]:
        values=[]
        for token in re.findall(r'\d[\d ,.]*\d|\d+', text):
            digits=re.sub(r'\D','',token)
            if len(digits) >= 5:
                values.append(f'{int(digits):,}'.replace(',', ' '))
        return values

    def parse_rows(text: str) -> dict[str,list[str]]:
        rows={}; current=None
        for line in text.splitlines():
            match=re.match(r'^\s*(4|5|6|7|8|9|10|11|12)\b(.*)$',line)
            if match:
                values=normalize_numbers(match.group(2))
                if len(values) >= 3:
                    current=match.group(1); rows[current]=values
                continue
            if current:
                rows[current].extend(normalize_numbers(line))
        return rows

    try:
        text=ocr_pdf(data) if data.startswith(b'%PDF') else extract_text('scan.png',data)
    except (ImportError, OSError, RuntimeError):
        return None
    rows=parse_rows(text)
    if rows:
        output=['| Echelon | Valeurs salariales extraites |','| --- | --- |']
        output.extend(f"| {echelon} | {' '.join(rows[echelon])} |" for echelon in sorted(rows,key=int))
        return '\n'.join(output)
    return None
